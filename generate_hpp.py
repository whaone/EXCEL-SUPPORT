#!/usr/bin/env python3
"""
Generator format Excel HPP (Harga Pokok Produksi) untuk usaha konveksi / sablon kaos.

Struktur workbook:
  1. Petunjuk      -> cara pakai
  2. Master        -> data harga kain per warna + biaya default (semua bisa diedit)
  3. Kalkulasi HPP -> input warna kain, design, jumlah; HPP & laba dihitung otomatis
  4. Ringkasan     -> rekap total + rekap per warna kain

Semua angka memakai FORMULA sehingga ketika master/biaya diubah, HPP ikut berubah
(fully customizable).
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, NamedStyle
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------- styling helpers
RP = '"Rp"#,##0'          # format rupiah
PCT = '0%'                # format persen
NUM = '#,##0'             # angka biasa

C_TITLE = "1F4E78"        # biru tua
C_HEAD = "2E75B6"         # biru header
C_SUB = "DDEBF7"          # biru muda
C_INPUT = "FFF2CC"        # kuning = sel input (boleh diedit)
C_CALC = "E2EFDA"         # hijau muda = sel hasil otomatis
C_MASTER = "FCE4D6"       # oranye muda = master
C_WHITE = "FFFFFF"

thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def style_title(cell):
    cell.font = Font(bold=True, size=14, color=C_WHITE)
    cell.fill = PatternFill("solid", fgColor=C_TITLE)
    cell.alignment = Alignment(horizontal="left", vertical="center")


def style_head(cell):
    cell.font = Font(bold=True, size=10, color=C_WHITE)
    cell.fill = PatternFill("solid", fgColor=C_HEAD)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = BORDER


def fill(cell, color):
    cell.fill = PatternFill("solid", fgColor=color)
    cell.border = BORDER


wb = Workbook()

# ================================================================ 1. PETUNJUK
ws = wb.active
ws.title = "Petunjuk"
ws.sheet_view.showGridLines = False
ws.column_dimensions["A"].width = 3
ws.column_dimensions["B"].width = 95

ws["B2"] = "FORMAT HPP KONVEKSI / SABLON KAOS"
style_title(ws["B2"])
ws.row_dimensions[2].height = 26

petunjuk = [
    ("", ""),
    ("APA INI?", "h"),
    ("File untuk menghitung Harga Pokok Produksi (HPP) kaos berdasarkan warna kain.", "p"),
    ("Satu warna kain bisa dipakai untuk beberapa design — tinggal tambah baris.", "p"),
    ("", ""),
    ("WARNA SEL (kode warna)", "h"),
    ("KUNING  = sel input, boleh kamu isi/ubah sesuai kebutuhan.", "p"),
    ("HIJAU   = hasil hitung otomatis (jangan diubah, sudah ada rumus).", "p"),
    ("ORANYE  = data master (harga kain & biaya default).", "p"),
    ("", ""),
    ("CARA PAKAI", "h"),
    ("1. Buka sheet 'Master'. Isi daftar warna kain, harga kain /Kg, dan", "p"),
    ("   estimasi berapa pcs kaos jadi dari 1 Kg kain. Isi juga biaya default.", "p"),
    ("2. Buka sheet 'Kalkulasi HPP'. Pilih Warna Kain dari dropdown,", "p"),
    ("   tulis nama Design, dan isi Jumlah (pcs).", "p"),
    ("3. Kolom harga kain, ongkos jahit, sablon akan terisi otomatis dari Master,", "p"),
    ("   tapi tetap boleh ditimpa manual per baris bila ada yang beda.", "p"),
    ("4. Isi ongkos transport & biaya lain-lain per batch bila ada.", "p"),
    ("5. HPP /pcs, harga jual, dan laba dihitung otomatis.", "p"),
    ("6. Lihat sheet 'Ringkasan' untuk total & rekap per warna kain.", "p"),
    ("", ""),
    ("RUMUS HPP /pcs", "h"),
    ("Biaya Kain /pcs + Ongkos Jahit + Ongkos Sablon + Biaya Lain /pcs", "p"),
    ("  + (Transport + Biaya Lain-lain) / Jumlah pcs", "p"),
    ("Harga Jual /pcs = HPP /pcs x (1 + Margin %)", "p"),
]
r = 4
for text, kind in petunjuk:
    c = ws.cell(row=r, column=2, value=text)
    if kind == "h":
        c.font = Font(bold=True, size=11, color=C_TITLE)
        c.fill = PatternFill("solid", fgColor=C_SUB)
    elif kind == "p":
        c.font = Font(size=10)
    r += 1

# legend kotak warna
for row, color, label in [(4, C_INPUT, "INPUT"), (4, None, None)]:
    pass

# ================================================================ 2. MASTER
ms = wb.create_sheet("Master")
ms.sheet_view.showGridLines = False
widths = {"A": 3, "B": 20, "C": 18, "D": 18, "E": 16, "F": 16}
for col, w in widths.items():
    ms.column_dimensions[col].width = w

ms["B2"] = "MASTER DATA — silakan edit angka di sini"
style_title(ms["B2"])
ms.merge_cells("B2:F2")
ms.row_dimensions[2].height = 24

# --- Tabel 1: Master Kain
ms["B4"] = "1. MASTER KAIN"
ms["B4"].font = Font(bold=True, size=11, color=C_TITLE)
ms["B4"].fill = PatternFill("solid", fgColor=C_SUB)
ms.merge_cells("B4:F4")

kain_headers = ["Warna Kain", "Jenis Kain", "Harga Kain /Kg", "Pcs per Kg"]
for i, h in enumerate(kain_headers):
    style_head(ms.cell(row=5, column=2 + i, value=h))

# data sample kain (boleh diedit / ditambah). Baris 6 - 25
kain_sample = [
    ("Hitam",   "Cotton Combed 30s", 95000, 4),
    ("Putih",   "Cotton Combed 30s", 90000, 4),
    ("Navy",    "Cotton Combed 30s", 98000, 4),
    ("Maroon",  "Cotton Combed 30s", 98000, 4),
    ("Abu Misty","Cotton Combed 30s", 92000, 4),
    ("Merah",   "Cotton Combed 24s", 99000, 3),
]
KAIN_FIRST, KAIN_LAST = 6, 25
for idx in range(KAIN_FIRST, KAIN_LAST + 1):
    vals = kain_sample[idx - KAIN_FIRST] if (idx - KAIN_FIRST) < len(kain_sample) else ("", "", None, None)
    for j in range(4):
        cell = ms.cell(row=idx, column=2 + j, value=vals[j] if vals[j] != "" else None)
        fill(cell, C_MASTER)
        if j == 2:
            cell.number_format = RP
        if j == 3:
            cell.number_format = NUM

# named range untuk VLOOKUP & dropdown
wb.defined_names.add  # noqa
from openpyxl.workbook.defined_name import DefinedName
wb.defined_names.add(DefinedName("MasterKain", attr_text=f"Master!$B${KAIN_FIRST}:$E${KAIN_LAST}"))
wb.defined_names.add(DefinedName("ListWarna", attr_text=f"Master!$B${KAIN_FIRST}:$B${KAIN_LAST}"))

# --- Tabel 2: Biaya default (parameter, pakai named range)
ms["B27"] = "2. BIAYA PRODUKSI DEFAULT (per pcs)"
ms["B27"].font = Font(bold=True, size=11, color=C_TITLE)
ms["B27"].fill = PatternFill("solid", fgColor=C_SUB)
ms.merge_cells("B27:F27")

params = [
    ("Ongkos Jahit /pcs", 8000, RP, "OngkosJahit"),
    ("Ongkos Sablon /pcs", 7000, RP, "OngkosSablon"),
    ("Biaya Lain /pcs (label, packaging, dll)", 3000, RP, "BiayaLain"),
    ("Margin Keuntungan", 0.40, PCT, "Margin"),
]
style_head(ms.cell(row=28, column=2, value="Komponen"))
style_head(ms.cell(row=28, column=3, value="Nilai Default"))
prow = 29
for label, val, fmt, name in params:
    lc = ms.cell(row=prow, column=2, value=label)
    fill(lc, C_SUB)
    lc.font = Font(size=10)
    vc = ms.cell(row=prow, column=3, value=val)
    fill(vc, C_MASTER)
    vc.number_format = fmt
    vc.alignment = Alignment(horizontal="right")
    wb.defined_names.add(DefinedName(name, attr_text=f"Master!$C${prow}"))
    prow += 1

# ================================================================ 3. KALKULASI HPP
ks = wb.create_sheet("Kalkulasi HPP")
ks.sheet_view.showGridLines = False

ks["A1"] = "KALKULASI HPP — KONVEKSI / SABLON KAOS"
style_title(ks["A1"])

# definisi kolom: (header, lebar, format, jenis)  jenis: input/calc
cols = [
    ("No", 5, NUM, "calc"),
    ("Warna Kain", 16, None, "input"),
    ("Design", 20, None, "input"),
    ("Jumlah (pcs)", 11, NUM, "input"),
    ("Harga Kain /Kg", 14, RP, "auto"),
    ("Pcs / Kg", 9, NUM, "auto"),
    ("Biaya Kain /pcs", 13, RP, "calc"),
    ("Ongkos Jahit /pcs", 12, RP, "auto"),
    ("Ongkos Sablon /pcs", 12, RP, "auto"),
    ("Biaya Lain /pcs", 12, RP, "auto"),
    ("Transport (batch)", 13, RP, "input"),
    ("Biaya Lain2 (batch)", 13, RP, "input"),
    ("HPP /pcs", 13, RP, "calc"),
    ("Total HPP", 15, RP, "calc"),
    ("Margin", 9, PCT, "auto"),
    ("Harga Jual /pcs", 14, RP, "calc"),
    ("Total Jual", 15, RP, "calc"),
    ("Laba", 15, RP, "calc"),
]
HEAD_ROW = 3
for i, (h, w, fmt, kind) in enumerate(cols):
    col = i + 1
    style_head(ks.cell(row=HEAD_ROW, column=col, value=h))
    ks.column_dimensions[get_column_letter(col)].width = w

ks.freeze_panes = "A4"

DATA_FIRST = HEAD_ROW + 1          # 4
N_ROWS = 40
DATA_LAST = DATA_FIRST + N_ROWS - 1

for row in range(DATA_FIRST, DATA_LAST + 1):
    r = row
    # A No
    ks.cell(row=r, column=1, value=f"=IF($B{r}=\"\",\"\",ROW()-{HEAD_ROW})")
    # B Warna (input) ; C Design (input) ; D Jumlah (input)
    # E Harga Kain /Kg auto from master (editable)
    ks.cell(row=r, column=5, value=f"=IF($B{r}=\"\",\"\",IFERROR(VLOOKUP($B{r},MasterKain,3,FALSE),0))")
    # F Pcs/Kg auto
    ks.cell(row=r, column=6, value=f"=IF($B{r}=\"\",\"\",IFERROR(VLOOKUP($B{r},MasterKain,4,FALSE),0))")
    # G Biaya Kain /pcs = E/F
    ks.cell(row=r, column=7, value=f"=IF(OR($B{r}=\"\",$F{r}=0),\"\",$E{r}/$F{r})")
    # H Ongkos Jahit auto
    ks.cell(row=r, column=8, value=f"=IF($B{r}=\"\",\"\",OngkosJahit)")
    # I Ongkos Sablon auto
    ks.cell(row=r, column=9, value=f"=IF($B{r}=\"\",\"\",OngkosSablon)")
    # J Biaya Lain /pcs auto
    ks.cell(row=r, column=10, value=f"=IF($B{r}=\"\",\"\",BiayaLain)")
    # K Transport batch (input) ; L Biaya lain2 batch (input)
    # M HPP /pcs
    ks.cell(row=r, column=13,
            value=(f"=IF(OR($B{r}=\"\",$D{r}=0),\"\","
                   f"$G{r}+$H{r}+$I{r}+$J{r}+(N($K{r})+N($L{r}))/$D{r})"))
    # N Total HPP = M*D
    ks.cell(row=r, column=14, value=f"=IF($M{r}=\"\",\"\",$M{r}*$D{r})")
    # O Margin auto
    ks.cell(row=r, column=15, value=f"=IF($B{r}=\"\",\"\",Margin)")
    # P Harga Jual /pcs = M*(1+O)
    ks.cell(row=r, column=16, value=f"=IF($M{r}=\"\",\"\",$M{r}*(1+$O{r}))")
    # Q Total Jual = P*D
    ks.cell(row=r, column=17, value=f"=IF($P{r}=\"\",\"\",$P{r}*$D{r})")
    # R Laba = Q-N
    ks.cell(row=r, column=18, value=f"=IF($Q{r}=\"\",\"\",$Q{r}-$N{r})")

    # apply format + warna per kolom
    for i, (h, w, fmt, kind) in enumerate(cols):
        c = ks.cell(row=r, column=i + 1)
        if fmt:
            c.number_format = fmt
        c.border = BORDER
        if kind == "input":
            c.fill = PatternFill("solid", fgColor=C_INPUT)
        elif kind == "auto":
            c.fill = PatternFill("solid", fgColor=C_INPUT)   # auto = terisi otomatis tapi boleh ditimpa
        elif kind == "calc":
            c.fill = PatternFill("solid", fgColor=C_CALC)

# baris TOTAL
tot = DATA_LAST + 1
ks.cell(row=tot, column=3, value="TOTAL").font = Font(bold=True)
for col, letter in [(4, "D"), (14, "N"), (17, "Q"), (18, "R")]:
    c = ks.cell(row=tot, column=col,
                value=f"=SUM({letter}{DATA_FIRST}:{letter}{DATA_LAST})")
    c.number_format = NUM if col == 4 else RP
    c.font = Font(bold=True, color=C_WHITE)
    c.fill = PatternFill("solid", fgColor=C_HEAD)
    c.border = BORDER
ks.cell(row=tot, column=3).fill = PatternFill("solid", fgColor=C_HEAD)
ks.cell(row=tot, column=3).font = Font(bold=True, color=C_WHITE)
ks.cell(row=tot, column=3).border = BORDER

# dropdown warna kain di kolom B
dv = DataValidation(type="list", formula1="=ListWarna", allow_blank=True)
dv.error = "Pilih warna dari daftar di sheet Master (atau tambahkan dulu di Master)."
dv.errorTitle = "Warna tidak ada"
dv.prompt = "Pilih warna kain"
ks.add_data_validation(dv)
dv.add(f"B{DATA_FIRST}:B{DATA_LAST}")

# contoh data terisi (3 baris) supaya langsung kelihatan hasilnya
ks["B4"], ks["C4"], ks["D4"] = "Hitam", "Logo Depan 1 Warna", 50
ks["B5"], ks["C5"], ks["D5"], ks["K5"] = "Hitam", "Full Print Belakang", 30, 25000
ks["B6"], ks["C6"], ks["D6"], ks["K6"] = "Putih", "Sablon 3 Warna", 40, 30000

# ================================================================ 4. RINGKASAN
rs = wb.create_sheet("Ringkasan")
rs.sheet_view.showGridLines = False
for col, w in {"A": 3, "B": 26, "C": 20, "D": 4, "E": 18, "F": 14, "G": 16, "H": 16}.items():
    rs.column_dimensions[col].width = w

rs["B2"] = "RINGKASAN HPP & LABA"
style_title(rs["B2"])
rs.merge_cells("B2:C2")

ringkas = [
    ("Total Produksi (pcs)", f"=SUM('Kalkulasi HPP'!D{DATA_FIRST}:D{DATA_LAST})", NUM),
    ("Total HPP", f"=SUM('Kalkulasi HPP'!N{DATA_FIRST}:N{DATA_LAST})", RP),
    ("Total Penjualan", f"=SUM('Kalkulasi HPP'!Q{DATA_FIRST}:Q{DATA_LAST})", RP),
    ("Total Laba", f"=SUM('Kalkulasi HPP'!R{DATA_FIRST}:R{DATA_LAST})", RP),
    ("Rata-rata HPP /pcs", f"=IFERROR(C5/C4,0)", RP),
    ("Margin Laba (%)", f"=IFERROR(C7/C6,0)", PCT),
]
rr = 4
for label, formula, fmt in ringkas:
    lc = rs.cell(row=rr, column=2, value=label)
    lc.font = Font(bold=True, size=10)
    fill(lc, C_SUB)
    vc = rs.cell(row=rr, column=3, value=formula)
    vc.number_format = fmt
    fill(vc, C_CALC)
    vc.alignment = Alignment(horizontal="right")
    vc.font = Font(bold=True)
    rr += 1

# rekap per warna kain (SUMIF)
rs["E2"] = "REKAP PER WARNA KAIN"
rs["E2"].font = Font(bold=True, size=11, color=C_TITLE)
rs["E2"].fill = PatternFill("solid", fgColor=C_SUB)
rs.merge_cells("E2:H2")

for i, h in enumerate(["Warna Kain", "Total pcs", "Total HPP", "Total Laba"]):
    style_head(rs.cell(row=3, column=5 + i, value=h))

warna_first = 4
for k in range(KAIN_FIRST, KAIN_LAST + 1):
    rr = warna_first + (k - KAIN_FIRST)
    # tampilkan warna hanya jika ada di master
    rs.cell(row=rr, column=5, value=f"=IF(Master!B{k}=\"\",\"\",Master!B{k})")
    rs.cell(row=rr, column=6,
            value=(f"=IF(Master!B{k}=\"\",\"\",SUMIF('Kalkulasi HPP'!$B${DATA_FIRST}:$B${DATA_LAST},"
                   f"Master!B{k},'Kalkulasi HPP'!$D${DATA_FIRST}:$D${DATA_LAST}))"))
    rs.cell(row=rr, column=7,
            value=(f"=IF(Master!B{k}=\"\",\"\",SUMIF('Kalkulasi HPP'!$B${DATA_FIRST}:$B${DATA_LAST},"
                   f"Master!B{k},'Kalkulasi HPP'!$N${DATA_FIRST}:$N${DATA_LAST}))"))
    rs.cell(row=rr, column=8,
            value=(f"=IF(Master!B{k}=\"\",\"\",SUMIF('Kalkulasi HPP'!$B${DATA_FIRST}:$B${DATA_LAST},"
                   f"Master!B{k},'Kalkulasi HPP'!$R${DATA_FIRST}:$R${DATA_LAST}))"))
    for col in range(5, 9):
        c = rs.cell(row=rr, column=col)
        c.border = BORDER
        if col == 6:
            c.number_format = NUM
        elif col in (7, 8):
            c.number_format = RP

OUT = "HPP_Konveksi_Kaos.xlsx"
wb.save(OUT)
print("Tersimpan:", OUT)
print("Sheets:", wb.sheetnames)
