#!/usr/bin/env python3
"""Generate a nicely formatted .xlsx production schedule for the t-shirt split."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

DAYS = ["Hari 1", "Hari 2", "Hari 3", "Hari 4", "Hari 5"]

# Per NORMAL design (multiplier 1). Design X uses multiplier 2.
NORMAL_ADULT = {
    ("S",   "short"): [1, 1, 1, 0, 0],
    ("S",   "long"):  [0, 0, 0, 1, 1],
    ("M",   "short"): [2, 2, 2, 2, 2],
    ("M",   "long"):  [2, 2, 2, 2, 2],
    ("L",   "short"): [3, 3, 3, 2, 2],
    ("L",   "long"):  [2, 2, 2, 3, 3],
    ("XL",  "short"): [2, 2, 2, 2, 2],
    ("XL",  "long"):  [2, 2, 2, 2, 2],
    ("2XL", "short"): [2, 2, 2, 2, 2],
    ("2XL", "long"):  [2, 2, 2, 2, 2],
    ("3XL", "short"): [1, 1, 1, 1, 1],
    ("3XL", "long"):  [1, 1, 1, 1, 1],
}
NORMAL_KIDS = {"S": [1, 1, 1, 1, 1], "M": [1, 1, 1, 1, 1], "L": [1, 1, 1, 1, 1]}

SIZES = ["S", "M", "L", "XL", "2XL", "3XL"]
VARIANTS = ["short", "long"]
KIDS_SIZES = ["S", "M", "L"]
NORMAL_DESIGNS = list("ABCDEFGHIJKLMN")
ALL_DESIGNS = [(d, 1) for d in NORMAL_DESIGNS] + [("X", 2)]

# ---- styles ----
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(bold=True, size=14, color="1F4E78")
DESIGN_FILL = PatternFill("solid", fgColor="DDEBF7")
X_FILL = PatternFill("solid", fgColor="FCE4D6")
TOTAL_FILL = PatternFill("solid", fgColor="FFF2CC")
TOTAL_FONT = Font(bold=True)
CENTER = Alignment(horizontal="center", vertical="center")
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = BORDER


def add(a, b):
    return [a[i] + b[i] for i in range(5)]


def add(a, b):
    return [a[i] + b[i] for i in range(5)]


def main():
    wb = openpyxl.Workbook()

    # ============ SHEET 1: Detail per design ============
    ws = wb.active
    ws.title = "Detail per Design"
    ws["A1"] = "PEMBAGIAN KAOS 5 HARI - 15 DESIGN (A-N + X)"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=10)

    header = ["Design", "Kategori", "Ukuran", "Variasi"] + DAYS + ["Total"]
    hr = 3
    ws.append([])
    for c, h in enumerate(header, 1):
        ws.cell(row=hr, column=c, value=h)
    style_header(ws, hr, len(header))
    ws.freeze_panes = "A4"

    grand_adult = {k: [0]*5 for k in NORMAL_ADULT}
    grand_kids = {k: [0]*5 for k in NORMAL_KIDS}
    r = hr + 1
    for design, mult in ALL_DESIGNS:
        start = r
        for size in SIZES:
            for var in VARIANTS:
                arr = [v*mult for v in NORMAL_ADULT[(size, var)]]
                ws.cell(row=r, column=1, value=design)
                ws.cell(row=r, column=2, value="Dewasa")
                ws.cell(row=r, column=3, value=size)
                ws.cell(row=r, column=4, value=var)
                for i, v in enumerate(arr):
                    ws.cell(row=r, column=5+i, value=v)
                ws.cell(row=r, column=10, value=sum(arr))
                grand_adult[(size, var)] = add(grand_adult[(size, var)], arr)
                r += 1
        for size in KIDS_SIZES:
            arr = [v*mult for v in NORMAL_KIDS[size]]
            ws.cell(row=r, column=1, value=design)
            ws.cell(row=r, column=2, value="Anak")
            ws.cell(row=r, column=3, value=size)
            ws.cell(row=r, column=4, value="-")
            for i, v in enumerate(arr):
                ws.cell(row=r, column=5+i, value=v)
            ws.cell(row=r, column=10, value=sum(arr))
            grand_kids[size] = add(grand_kids[size], arr)
            r += 1
        # design subtotal row
        ws.cell(row=r, column=1, value=design)
        ws.cell(row=r, column=2, value="TOTAL design")
        sub = [0]*5
        for rr in range(start, r):
            for i in range(5):
                sub[i] += ws.cell(row=rr, column=5+i).value
        for i, v in enumerate(sub):
            ws.cell(row=r, column=5+i, value=v)
        ws.cell(row=r, column=10, value=sum(sub))
        for c in range(1, 11):
            ws.cell(row=r, column=c).fill = TOTAL_FILL
            ws.cell(row=r, column=c).font = TOTAL_FONT
        # color the design block first column
        fill = X_FILL if design == "X" else DESIGN_FILL
        for rr in range(start, r):
            ws.cell(row=rr, column=1).fill = fill
        r += 1

    # borders + alignment for all data cells
    for rr in range(hr, r):
        for c in range(1, 11):
            cell = ws.cell(row=rr, column=c)
            cell.border = BORDER
            if c >= 3:
                cell.alignment = CENTER

    # ============ SHEET 2: Total per hari ============
    ws2 = wb.create_sheet("Total per Hari")
    ws2["A1"] = "TOTAL GABUNGAN SEMUA 15 DESIGN PER HARI"
    ws2["A1"].font = TITLE_FONT
    ws2.merge_cells(start_row=1, start_column=1, end_row=1, end_column=9)
    header2 = ["Kategori", "Ukuran", "Variasi"] + DAYS + ["Total"]
    for c, h in enumerate(header2, 1):
        ws2.cell(row=3, column=c, value=h)
    style_header(ws2, 3, len(header2))
    r2 = 4
    for size in SIZES:
        for var in VARIANTS:
            arr = grand_adult[(size, var)]
            ws2.cell(row=r2, column=1, value="Dewasa")
            ws2.cell(row=r2, column=2, value=size)
            ws2.cell(row=r2, column=3, value=var)
            for i, v in enumerate(arr):
                ws2.cell(row=r2, column=4+i, value=v)
            ws2.cell(row=r2, column=9, value=sum(arr))
            r2 += 1
    for size in KIDS_SIZES:
        arr = grand_kids[size]
        ws2.cell(row=r2, column=1, value="Anak")
        ws2.cell(row=r2, column=2, value=size)
        ws2.cell(row=r2, column=3, value="-")
        for i, v in enumerate(arr):
            ws2.cell(row=r2, column=4+i, value=v)
        ws2.cell(row=r2, column=9, value=sum(arr))
        r2 += 1
    # grand total per day
    day_tot = [0]*5
    for k in grand_adult:
        day_tot = add(day_tot, grand_adult[k])
    for k in grand_kids:
        day_tot = add(day_tot, grand_kids[k])
    ws2.cell(row=r2, column=1, value="TOTAL")
    ws2.cell(row=r2, column=2, value="SEMUA")
    ws2.cell(row=r2, column=3, value="-")
    for i, v in enumerate(day_tot):
        ws2.cell(row=r2, column=4+i, value=v)
    ws2.cell(row=r2, column=9, value=sum(day_tot))
    for c in range(1, 10):
        ws2.cell(row=r2, column=c).fill = TOTAL_FILL
        ws2.cell(row=r2, column=c).font = TOTAL_FONT
    for rr in range(3, r2 + 1):
        for c in range(1, 10):
            cell = ws2.cell(row=rr, column=c)
            cell.border = BORDER
            if c >= 2:
                cell.alignment = CENTER

    # ============ SHEET 3: Ringkasan ============
    ws3 = wb.create_sheet("Ringkasan")
    ws3["A1"] = "RINGKASAN PEMBAGIAN KAOS"
    ws3["A1"].font = TITLE_FONT
    info = [
        ["", ""],
        ["Jumlah design", "15 (A-N normal + X spesial)"],
        ["Design normal (A-N)", "100 pcs dewasa + 15 pcs anak / design"],
        ["Design X", "200 pcs dewasa + 30 pcs anak"],
        ["Total dewasa", 1600],
        ["Total anak", 240],
        ["GRAND TOTAL", 1840],
        ["Produksi per hari", 368],
        ["", ""],
        ["Catatan ukuran S", "stok 5/design -> short di Hari 1-3, long di Hari 4-5"],
        ["Catatan ukuran L", "13 short vs 12 long -> 3/3/3/2/2 dan 2/2/2/3/3"],
        ["Ukuran lain", "short & long terbagi rata tiap hari"],
    ]
    rr = 3
    for k, v in info:
        ws3.cell(row=rr, column=1, value=k).font = Font(bold=True)
        ws3.cell(row=rr, column=2, value=v)
        rr += 1
    ws3.column_dimensions["A"].width = 22
    ws3.column_dimensions["B"].width = 55

    # column widths for sheet1 & sheet2
    for w, col in zip([10, 14, 8, 9, 9, 9, 9, 9, 9, 9], range(1, 11)):
        ws.column_dimensions[get_column_letter(col)].width = w
    for w, col in zip([10, 8, 9, 9, 9, 9, 9, 9, 9], range(1, 10)):
        ws2.column_dimensions[get_column_letter(col)].width = w

    out = "pembagian_kaos_5hari.xlsx"
    wb.save(out)
    print("Saved:", out)
    print("Total per hari:", day_tot, "Grand total:", sum(day_tot))


if __name__ == "__main__":
    main()
